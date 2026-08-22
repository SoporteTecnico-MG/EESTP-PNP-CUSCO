from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from django.core.management.base import BaseCommand

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

DIAS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]
ESPECIALIDADES = ["INVESTIGACION CRIMINAL", "ORDEN PUBLICO Y SEGURIDAD CIUDADANA"]


def _hoja(wb, nombre, headers, ejemplo, anchos=None):
    ws = wb.create_sheet(nombre)
    for col, titulo in enumerate(headers, start=1):
        celda = ws.cell(row=1, column=col, value=titulo)
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT
    for fila_idx, fila in enumerate(ejemplo, start=2):
        for col, valor in enumerate(fila, start=1):
            ws.cell(row=fila_idx, column=col, value=valor)
    if anchos:
        for col, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = ancho
    ws.freeze_panes = "A2"
    return ws


class Command(BaseCommand):
    help = "Genera la plantilla Excel (Docentes, Periodos, Aulas, Cursos, Horarios) para la carga masiva de datos."

    def add_arguments(self, parser):
        parser.add_argument(
            "--salida",
            default="plantilla_control_docentes.xlsx",
            help="Ruta del archivo .xlsx a generar.",
        )

    def handle(self, *args, **options):
        wb = Workbook()
        wb.remove(wb.active)

        # --- Hoja Instrucciones ---
        ws_instr = wb.create_sheet("Instrucciones")
        instrucciones = [
            "PLANTILLA DE CARGA — CONTROL DE ASISTENCIA DOCENTE EESTP PNP CUSCO",
            "",
            "Completa las 5 hojas EN ESTE ORDEN: Docentes, Periodos, Aulas, Cursos, Horarios.",
            "No cambies los nombres de las columnas (fila 1) ni el orden de las hojas.",
            "",
            "HOJA 'Docentes':",
            "- ID Biometrico: el número (correlativo) tal como aparece enrolado en el equipo ZKTeco.",
            "  Este es el único dato realmente obligatorio junto con el nombre — el control de asistencia",
            "  funciona con el ID Biometrico, no con el DNI.",
            "- DNI: puedes dejarlo vacío si todavía no lo tienes a mano y regularizarlo después; mientras",
            "  esté vacío, en la hoja Horarios usa el ID Biometrico en vez del DNI para ese docente.",
            "- Fecha Ingreso: formato DD/MM/AAAA (opcional).",
            "- Estado: ACTIVO o INACTIVO (si lo dejas vacío, se asume ACTIVO).",
            "",
            "HOJA 'Periodos':",
            "- Una fila por cada período académico de cada promoción (normalmente hasta 6 filas por promoción).",
            "- Promocion: nombre exacto, ej. 'JUSTICIEROS 2025-I' (debe escribirse igual en todas las hojas).",
            "- Fechas en formato DD/MM/AAAA.",
            "",
            "HOJA 'Aulas':",
            "- Una fila por cada aula de cada promoción (hasta 14 por promoción).",
            "- Aula: número (1, 2, 3...) o formato A_1, B_2, C_3... — el código se genera automático.",
            "- Especialidad: SOLO llenar desde el III período en adelante. Déjalo vacío en I y II período,",
            "  ya que ahí todavía no hay división por especialidad.",
            "- IMPORTANTE: la especialidad de un aula se define UNA SOLA VEZ aquí, no se repite ni se",
            "  puede contradecir en la hoja Horarios — evita el error de poner un aula OP con un curso IC.",
            "",
            "HOJA 'Cursos':",
            "- Una fila por cada curso que existe (sin repetir), independiente de en qué aula se dicte.",
            "- Especialidad: vacío si el curso es común a ambas especialidades (se dicta igual en todas",
            "  las aulas); si es un curso propio de una sola especialidad, indícala aquí UNA SOLA VEZ.",
            "",
            "HOJA 'Horarios':",
            "- Una fila por cada bloque de clase: un docente, en un aula, dictando un curso, un día específico.",
            "- Aula y Curso deben existir ya en sus hojas correspondientes (no repitas la especialidad aquí,",
            "  el sistema la toma automáticamente de las hojas Aulas/Cursos y avisa si hay una contradicción,",
            "  por ejemplo un curso propio de Investigación Criminal puesto en una aula de Orden Público).",
            "- Si un curso dura varias horas pedagógicas seguidas (ej. 4 horas), pon una sola fila con",
            "  Hora Inicio = inicio del bloque y Hora Fin = fin del bloque (ej. 08:00 y 11:10).",
            "- DNI o ID Biometrico Docente: puedes poner el DNI, o si todavía no lo tienes a mano, el",
            "  ID Biometrico del docente (el mismo que pusiste en la hoja Docentes). El docente debe",
            "  existir ya en la hoja 'Docentes' con ese mismo dato.",
            "- Hora Inicio / Hora Fin: en formato HH:MM (24 horas), deben coincidir con el inicio/fin real",
            "  de alguna hora pedagógica de la grilla (08:00, 08:45, 09:40, 10:25, 11:20, 12:05, 15:00,",
            "  15:45, 16:40, 17:25 como inicios; 08:45, 09:30, 10:25, 11:10, 12:05, 12:50, 15:45, 16:30,",
            "  17:25, 18:10 como fines).",
            "",
            "Una vez lleno, envíame el archivo y lo importo con el comando import_excel.",
        ]
        for i, linea in enumerate(instrucciones, start=1):
            ws_instr.cell(row=i, column=1, value=linea)
        ws_instr.cell(row=1, column=1).font = Font(bold=True, size=13)
        ws_instr.column_dimensions["A"].width = 110

        # --- Hoja Docentes ---
        _hoja(
            wb,
            "Docentes",
            ["DNI", "Apellidos y Nombres", "Grado", "Celular", "ID Biometrico", "Fecha Ingreso", "Estado"],
            [
                ["70123456", "PEREZ GOMEZ JUAN CARLOS", "MAYOR PNP", "987654321", "45", "15/03/2020", "ACTIVO"],
                ["", "QUISPE MAMANI ROSA ELENA", "CAPITAN PNP", "", "112", "", "ACTIVO"],
            ],
            anchos=[12, 32, 16, 14, 14, 14, 10],
        )

        # --- Hoja Periodos ---
        _hoja(
            wb,
            "Periodos",
            ["Promocion", "Anio Ingreso", "Fecha Inicio Formacion", "Numero Periodo", "Nombre Periodo", "Fecha Inicio", "Fecha Fin"],
            [
                ["JUSTICIEROS 2025-I", 2025, "01/03/2025", 3, "III Periodo", "01/07/2026", "31/10/2026"],
                ["AGUERRIDOS 2026-I", 2026, "01/03/2026", 1, "I Periodo", "01/03/2026", "30/06/2026"],
            ],
            anchos=[24, 12, 20, 14, 18, 14, 14],
        )

        # --- Hoja Aulas ---
        _hoja(
            wb,
            "Aulas",
            ["Promocion", "Aula", "Especialidad"],
            [
                ["JUSTICIEROS 2025-I", "A_1", "ORDEN PUBLICO Y SEGURIDAD CIUDADANA"],
                ["JUSTICIEROS 2025-I", "K_11", "INVESTIGACION CRIMINAL"],
                ["AGUERRIDOS 2026-I", "A_1", ""],
            ],
            anchos=[24, 10, 32],
        )

        # --- Hoja Cursos ---
        _hoja(
            wb,
            "Cursos",
            ["Curso", "Especialidad"],
            [
                ["DERECHOS HUMANOS", ""],
                ["INVESTIGACION DEL DELITO", "INVESTIGACION CRIMINAL"],
                ["ORDEN PUBLICO Y SEGURIDAD VIAL", "ORDEN PUBLICO Y SEGURIDAD CIUDADANA"],
            ],
            anchos=[32, 32],
        )

        # --- Hoja Horarios ---
        _hoja(
            wb,
            "Horarios",
            [
                "Promocion", "Numero Periodo", "Aula", "Curso", "DNI o ID Biometrico Docente",
                "Horas Pedagogicas", "Dia", "Hora Inicio", "Hora Fin",
            ],
            [
                [
                    "JUSTICIEROS 2025-I", 3, "A_1", "ORDEN PUBLICO Y SEGURIDAD VIAL",
                    "70123456", 4, "LUNES", "08:00", "11:10",
                ],
                [
                    "JUSTICIEROS 2025-I", 3, "A_1", "DERECHOS HUMANOS",
                    "45", 3, "MARTES", "08:00", "10:25",
                ],
            ],
            anchos=[24, 14, 8, 32, 20, 16, 12, 12, 12],
        )

        # Validaciones simples de lista desplegable
        ws_aulas = wb["Aulas"]
        dv_esp_aula = DataValidation(type="list", formula1=f'"{",".join(ESPECIALIDADES)}"', allow_blank=True)
        ws_aulas.add_data_validation(dv_esp_aula)
        dv_esp_aula.add("C2:C500")

        ws_cursos = wb["Cursos"]
        dv_esp_curso = DataValidation(type="list", formula1=f'"{",".join(ESPECIALIDADES)}"', allow_blank=True)
        ws_cursos.add_data_validation(dv_esp_curso)
        dv_esp_curso.add("B2:B500")

        ws_horarios = wb["Horarios"]
        dv_dia = DataValidation(type="list", formula1=f'"{",".join(DIAS)}"', allow_blank=False)
        ws_horarios.add_data_validation(dv_dia)
        dv_dia.add("G2:G500")

        ws_docentes = wb["Docentes"]
        dv_estado = DataValidation(type="list", formula1='"ACTIVO,INACTIVO"', allow_blank=True)
        ws_docentes.add_data_validation(dv_estado)
        dv_estado.add("G2:G500")

        salida = options["salida"]
        wb.save(salida)
        self.stdout.write(self.style.SUCCESS(f"Plantilla generada en: {salida}"))

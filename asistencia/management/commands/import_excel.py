import datetime
import re

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from asistencia.models import (
    Aula,
    Asignacion,
    BloqueHorario,
    Curso,
    Docente,
    Especialidad,
    HoraPedagogica,
    OfertaCurso,
    PeriodoAcademico,
    Promocion,
)

DIA_A_NUMERO = {
    "LUNES": 1,
    "MARTES": 2,
    "MIERCOLES": 3,
    "MIÉRCOLES": 3,
    "JUEVES": 4,
    "VIERNES": 5,
    "SABADO": 6,
    "SÁBADO": 6,
    "DOMINGO": 7,
}


ERRORES_EXCEL = {"#N/A", "#N/D", "N/D", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}


def _str(v):
    if v is None:
        return ""
    texto = str(v).strip()
    if texto.upper() in ERRORES_EXCEL:
        return ""
    return texto


def _numero_aula(v):
    """Acepta '1', 'A_1', 'B_2', 'Aula 3', etc. y devuelve el número entero."""
    match = re.search(r"\d+", str(v))
    if not match:
        raise ValueError(f"No se pudo obtener el número de aula de '{v}'.")
    return int(match.group())


def _fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if not v:
        return None
    return datetime.datetime.strptime(str(v).strip(), "%d/%m/%Y").date()


def _hora(v):
    if isinstance(v, datetime.time):
        return v
    if isinstance(v, datetime.datetime):
        return v.time()
    return datetime.datetime.strptime(str(v).strip(), "%H:%M").time()


class Command(BaseCommand):
    help = (
        "Importa docentes, períodos, aulas, cursos y horarios desde la plantilla Excel "
        "(hojas Docentes, Periodos, Aulas, Cursos, Horarios)."
    )

    def add_arguments(self, parser):
        parser.add_argument("archivo", help="Ruta al archivo .xlsx lleno.")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Valida y muestra lo que haría, sin escribir en la base de datos.",
        )

    def handle(self, *args, **options):
        ruta = options["archivo"]
        dry_run = options["dry_run"]

        try:
            wb = load_workbook(ruta, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"No se encontró el archivo: {ruta}")

        for hoja in ("Docentes", "Periodos", "Aulas", "Cursos", "Horarios"):
            if hoja not in wb.sheetnames:
                raise CommandError(f"Falta la hoja '{hoja}' en el archivo.")

        errores = []
        resumen = {
            "docentes": 0, "periodos": 0, "aulas": 0, "cursos": 0,
            "asignaciones": 0, "bloques": 0,
        }

        with transaction.atomic():
            self._importar_docentes(wb["Docentes"], errores, resumen)
            self._importar_periodos(wb["Periodos"], errores, resumen)
            self._importar_aulas(wb["Aulas"], errores, resumen)
            self._importar_cursos(wb["Cursos"], errores, resumen)
            self._importar_horarios(wb["Horarios"], errores, resumen)

            if errores or dry_run:
                transaction.set_rollback(True)

        if errores:
            self.stderr.write(self.style.ERROR(f"Se encontraron {len(errores)} error(es). No se guardó nada:"))
            for e in errores:
                self.stderr.write(f"  - {e}")
            return

        if dry_run:
            self.stdout.write(self.style.WARNING("Dry-run: nada se guardó en la base de datos. Esto es lo que se habría hecho:"))
        else:
            self.stdout.write(self.style.SUCCESS("Importación procesada:"))
        for clave, valor in resumen.items():
            self.stdout.write(f"  {clave}: {valor}")

    # ------------------------------------------------------------------

    def _importar_docentes(self, ws, errores, resumen):
        for fila in ws.iter_rows(min_row=2, values_only=False):
            valores = [c.value for c in fila]
            if not any(valores):
                continue
            dni, nombres, grado, celular, id_bio, fecha_ing, estado = (valores + [None] * 7)[:7]
            fila_num = fila[0].row
            if not nombres or not id_bio:
                errores.append(f"Docentes fila {fila_num}: Apellidos y Nombres, e ID Biometrico son obligatorios (el DNI puede quedar vacío por ahora).")
                continue
            try:
                fecha_ingreso = _fecha(fecha_ing) if fecha_ing else None
            except ValueError:
                errores.append(f"Docentes fila {fila_num}: fecha de ingreso inválida ('{fecha_ing}'), usa DD/MM/AAAA.")
                continue

            Docente.objects.update_or_create(
                id_biometrico=_str(id_bio),
                defaults={
                    "dni": _str(dni) or None,
                    "apellidos_nombres": _str(nombres),
                    "grado": _str(grado),
                    "celular": _str(celular),
                    "fecha_ingreso": fecha_ingreso,
                    "estado": (_str(estado) or "ACTIVO").upper(),
                },
            )
            resumen["docentes"] += 1

    def _importar_periodos(self, ws, errores, resumen):
        for fila in ws.iter_rows(min_row=2, values_only=False):
            valores = [c.value for c in fila]
            if not any(valores):
                continue
            (
                promocion_nombre, anio_ingreso, fecha_inicio_form,
                num_periodo, nombre_periodo, fecha_inicio, fecha_fin,
            ) = (valores + [None] * 7)[:7]
            fila_num = fila[0].row
            if not promocion_nombre or not num_periodo:
                errores.append(f"Periodos fila {fila_num}: Promocion y Numero Periodo son obligatorios.")
                continue
            try:
                promocion, _ = Promocion.objects.update_or_create(
                    nombre=_str(promocion_nombre),
                    defaults={
                        "anio_ingreso": int(anio_ingreso) if anio_ingreso else datetime.date.today().year,
                        "fecha_inicio_formacion": _fecha(fecha_inicio_form) or datetime.date.today(),
                    },
                )
                PeriodoAcademico.objects.update_or_create(
                    promocion=promocion,
                    numero_periodo=int(num_periodo),
                    defaults={
                        "nombre": _str(nombre_periodo),
                        "fecha_inicio": _fecha(fecha_inicio),
                        "fecha_fin": _fecha(fecha_fin),
                    },
                )
            except (ValueError, TypeError) as e:
                errores.append(f"Periodos fila {fila_num}: dato inválido ({e}).")
                continue
            resumen["periodos"] += 1

    def _importar_aulas(self, ws, errores, resumen):
        for fila in ws.iter_rows(min_row=2, values_only=False):
            valores = [c.value for c in fila]
            if not any(valores):
                continue
            promocion_nombre, num_aula, especialidad_nombre = (valores + [None] * 3)[:3]
            fila_num = fila[0].row
            if not promocion_nombre or not num_aula:
                errores.append(f"Aulas fila {fila_num}: Promocion y Aula son obligatorios.")
                continue

            try:
                promocion = Promocion.objects.get(nombre=_str(promocion_nombre))
            except Promocion.DoesNotExist:
                errores.append(f"Aulas fila {fila_num}: la promoción '{promocion_nombre}' no existe (revisa la hoja Periodos).")
                continue

            try:
                numero = _numero_aula(num_aula)
            except ValueError as e:
                errores.append(f"Aulas fila {fila_num}: {e}")
                continue

            especialidad = None
            if especialidad_nombre:
                especialidad, _ = Especialidad.objects.get_or_create(nombre=_str(especialidad_nombre).upper())

            Aula.objects.update_or_create(
                promocion=promocion,
                numero=numero,
                defaults={"especialidad": especialidad},
            )
            resumen["aulas"] += 1

    def _importar_cursos(self, ws, errores, resumen):
        for fila in ws.iter_rows(min_row=2, values_only=False):
            valores = [c.value for c in fila]
            if not any(valores):
                continue
            curso_nombre, especialidad_nombre = (valores + [None] * 2)[:2]
            fila_num = fila[0].row
            if not curso_nombre:
                errores.append(f"Cursos fila {fila_num}: Curso es obligatorio.")
                continue

            especialidad = None
            if especialidad_nombre:
                especialidad, _ = Especialidad.objects.get_or_create(nombre=_str(especialidad_nombre).upper())

            Curso.objects.update_or_create(
                nombre=_str(curso_nombre),
                defaults={"especialidad": especialidad},
            )
            resumen["cursos"] += 1

    def _importar_horarios(self, ws, errores, resumen):
        for fila in ws.iter_rows(min_row=2, values_only=False):
            valores = [c.value for c in fila]
            if not any(valores):
                continue
            (
                promocion_nombre, num_periodo, num_aula, curso_nombre,
                docente_ref, horas_pedagogicas, dia, hora_ini, hora_fin,
            ) = (valores + [None] * 9)[:9]
            fila_num = fila[0].row

            faltantes = [
                nombre for nombre, val in [
                    ("Promocion", promocion_nombre), ("Numero Periodo", num_periodo),
                    ("Aula", num_aula), ("Curso", curso_nombre),
                    ("DNI o ID Biometrico Docente", docente_ref),
                    ("Dia", dia), ("Hora Inicio", hora_ini), ("Hora Fin", hora_fin),
                ] if not val
            ]
            if faltantes:
                errores.append(f"Horarios fila {fila_num}: faltan campos obligatorios: {', '.join(faltantes)}.")
                continue

            try:
                promocion = Promocion.objects.get(nombre=_str(promocion_nombre))
            except Promocion.DoesNotExist:
                errores.append(f"Horarios fila {fila_num}: la promoción '{promocion_nombre}' no existe (revisa la hoja Periodos).")
                continue

            try:
                periodo = PeriodoAcademico.objects.get(promocion=promocion, numero_periodo=int(num_periodo))
            except PeriodoAcademico.DoesNotExist:
                errores.append(f"Horarios fila {fila_num}: no existe el período {num_periodo} para '{promocion_nombre}'.")
                continue

            docente_ref_str = _str(docente_ref)
            docente = Docente.objects.filter(dni=docente_ref_str).first() or Docente.objects.filter(
                id_biometrico=docente_ref_str
            ).first()
            if docente is None:
                errores.append(
                    f"Horarios fila {fila_num}: no se encontró ningún docente con DNI o ID Biometrico "
                    f"'{docente_ref}' (revisa la hoja Docentes)."
                )
                continue

            try:
                numero_aula = _numero_aula(num_aula)
            except ValueError as e:
                errores.append(f"Horarios fila {fila_num}: {e}")
                continue

            try:
                aula = Aula.objects.get(promocion=promocion, numero=numero_aula)
            except Aula.DoesNotExist:
                errores.append(
                    f"Horarios fila {fila_num}: el aula '{num_aula}' de '{promocion_nombre}' no existe "
                    f"(revisa la hoja Aulas)."
                )
                continue

            try:
                curso = Curso.objects.get(nombre=_str(curso_nombre))
            except Curso.DoesNotExist:
                errores.append(
                    f"Horarios fila {fila_num}: el curso '{curso_nombre}' no existe (revisa la hoja Cursos)."
                )
                continue

            # Regla de negocio: un curso propio de una especialidad solo puede dictarse en un
            # aula de esa misma especialidad. Un curso común (sin especialidad) va en cualquier aula.
            if curso.especialidad_id is not None and aula.especialidad_id is not None:
                if curso.especialidad_id != aula.especialidad_id:
                    errores.append(
                        f"Horarios fila {fila_num}: el curso '{curso.nombre}' es de "
                        f"'{curso.especialidad.nombre}' pero el aula '{aula.codigo}' es de "
                        f"'{aula.especialidad.nombre}' — no coinciden."
                    )
                    continue
            elif curso.especialidad_id is not None and aula.especialidad_id is None:
                errores.append(
                    f"Horarios fila {fila_num}: el curso '{curso.nombre}' es propio de "
                    f"'{curso.especialidad.nombre}' pero el aula '{aula.codigo}' todavía no tiene "
                    f"especialidad asignada (revisa la hoja Aulas)."
                )
                continue

            oferta, _ = OfertaCurso.objects.get_or_create(periodo_academico=periodo, curso=curso)

            Asignacion.objects.update_or_create(
                aula=aula,
                oferta_curso=oferta,
                defaults={"docente": docente},
            )
            resumen["asignaciones"] += 1

            dia_num = DIA_A_NUMERO.get(_str(dia).upper())
            if not dia_num:
                errores.append(f"Horarios fila {fila_num}: día inválido '{dia}'.")
                continue

            try:
                t_ini = _hora(hora_ini)
                t_fin = _hora(hora_fin)
            except ValueError:
                errores.append(f"Horarios fila {fila_num}: hora inválida, usa HH:MM.")
                continue

            try:
                hp_inicio = HoraPedagogica.objects.get(hora_inicio=t_ini, tipo="CLASE")
                hp_fin = HoraPedagogica.objects.get(hora_fin=t_fin, tipo="CLASE")
            except HoraPedagogica.DoesNotExist:
                errores.append(
                    f"Horarios fila {fila_num}: Hora Inicio/Fin ({hora_ini}/{hora_fin}) no coincide "
                    f"con ningún bloque de la grilla de horas pedagógicas."
                )
                continue

            # El horario de un curso es el mismo para todas las aulas que lo dicten en el
            # período: se fija una sola vez. Si ya existe con otro día/hora, es un conflicto
            # real en los datos (alguna fila del Excel no coincide con las demás).
            bloque_existente = BloqueHorario.objects.filter(oferta_curso=oferta).first()
            if bloque_existente and (
                bloque_existente.dia_semana != dia_num
                or bloque_existente.hora_pedagogica_inicio_id != hp_inicio.id
                or bloque_existente.hora_pedagogica_fin_id != hp_fin.id
            ):
                errores.append(
                    f"Horarios fila {fila_num}: el curso '{curso.nombre}' en '{periodo}' ya tiene "
                    f"horario fijado ({bloque_existente}) distinto al de esta fila — el horario debe "
                    f"ser el mismo para todas las aulas. Corrige la fila para que coincida."
                )
                continue

            BloqueHorario.objects.get_or_create(
                oferta_curso=oferta,
                dia_semana=dia_num,
                hora_pedagogica_inicio=hp_inicio,
                hora_pedagogica_fin=hp_fin,
            )
            resumen["bloques"] += 1
